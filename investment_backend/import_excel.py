#!/usr/bin/env python3
"""
Excel Import Script for Investments App
Reads from /home/armin/finances/NPV(Armin).xlsx and populates the database.
"""
import os
import sys
import openpyxl
from datetime import datetime
import psycopg2

# Database configuration
POSTGRES_HOST = os.getenv("POSTGRES_HOST", "ThinkTank")
POSTGRES_USER = os.getenv("POSTGRES_USER", "postgres")
POSTGRES_PASSWORD = os.getenv("POSTGRES_PASSWORD", "changeme")
POSTGRES_PORT = int(os.getenv("POSTGRES_PORT", 5432))
INVESTMENTS_DB = os.getenv("INVESTMENTS_DB", "Investments")
EXCEL_FILE = "/home/armin/finances/NPV(Armin).xlsx"

# Sheet mappings
# (SheetName, Institution, FundName, InvestmentType, Ticker, Currency)
INVESTMENT_SHEETS = [
    ("Allan Gray", "Allan Gray", "Allan Gray Equity Fund", "Unit Trust", "AGEF", "ZAR"),
    ("Allan Gray Balanced", "Allan Gray", "Allan Gray Balanced Fund", "Unit Trust", "AGBF", "ZAR"),
    ("TFSA Coronation", "Coronation", "Top 20 Fund (TFI A class)", "TFSA Unit Trust", "COR-TOP20", "ZAR"),
    ("AF_ISFAC", "Alexander Forbes", "AF Flexible Fund of Funds A (ISFAC)", "Unit Trust", "AF-ISFAC", "ZAR"),
    ("AF_ISGE", "Alexander Forbes", "AF Global Equity Feeder Fund A (ISGE)", "Unit Trust", "AF-ISGE", "ZAR"),
    ("Voluntary_PSG", "PSG Wealth", "Voluntary Investment (202402190101)", "Unit Trust", "PSG-VOL", "ZAR"),
    ("PSG_E", "PSG Wealth", "PSG Global Equity Feeder Fund (E)", "Unit Trust", "PSG-GE-E", "ZAR"),
    ("PSG_C", "PSG Wealth", "PSG Global Flexible Feeder Fund (C)", "Unit Trust", "PSG-GF-C", "ZAR"),
    ("PSG_TFSA", "PSG Wealth", "Tax Free Investment (202410100076)", "TFSA Unit Trust", "PSG-TFSA", "ZAR"),
    ("PSG_RA", "PSG Wealth", "PSG Balanced Fund (E)", "RA Unit Trust", "PSG-BAL-E", "ZAR"),
    ("EE_ZAR", "Easy Equities", "Satrix Nasdaq 100 ETF (STXNDQ) ZAR", "ETF", "STXNDQ.JO", "ZAR"),
    ("EE_TFSA_S", "Easy Equities", "Satrix Nasdaq 100 ETF (STXNDQ) TFSA", "ETF (TFSA)", "STXNDQ.JO", "ZAR"),
    ("EE_TFSA", "Easy Equities", "Satrix MSCI World (STXWDM) TFSA", "ETF (TFSA)", "STXWDM.JO", "ZAR"),
    ("EE_USD", "Easy Equities", "iShares MSCI World ETF (URTH) USD", "ETF", "URTH", "USD"),
    ("EE_GBP", "Easy Equities", "iShares Core EURO STOXX 50 UCITS ETF EUR (Dist EUE.L)", "ETF", "EUE.L", "GBP"),
]

def get_connection():
    return psycopg2.connect(
        host=POSTGRES_HOST,
        database=INVESTMENTS_DB,
        user=POSTGRES_USER,
        password=POSTGRES_PASSWORD,
        port=POSTGRES_PORT
    )

def clear_existing_data(cursor):
    """Optional: clear specific tables before import to avoid duplicates."""
    print("Clearing existing transactions and investments...")
    cursor.execute("DELETE FROM transactions")
    cursor.execute("DELETE FROM unit_prices")
    cursor.execute("DELETE FROM returns")
    cursor.execute("DELETE FROM fees")
    cursor.execute("DELETE FROM tax")
    cursor.execute("DELETE FROM dividends")
    cursor.execute("DELETE FROM investment_metrics")
    cursor.execute("DELETE FROM investment_source_meta")
    cursor.execute("DELETE FROM investments")

def import_excel_data():
    if not os.path.exists(EXCEL_FILE):
        print(f"Error: {EXCEL_FILE} not found.")
        sys.exit(1)
        
    print(f"Loading Excel file: {EXCEL_FILE}")
    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    
    conn = get_connection()
    cursor = conn.cursor()
    
    try:
        # Optional: uncomment to clear existing data before import
        # clear_existing_data(cursor)
        
        for sheet_name, inst, fund, inv_type, ticker, curr in INVESTMENT_SHEETS:
            if sheet_name not in wb.sheetnames:
                print(f"⚠️ Sheet '{sheet_name}' not found, skipping...")
                continue
                
            print(f"Processing sheet: {sheet_name} -> {fund}")
            ws = wb[sheet_name]
            
            # Find all transactions (Date and Value)
            # Assuming Date is in Col A, Value in Col B
            transactions = []
            for row in ws.iter_rows(min_row=2, max_col=2, values_only=True):
                if not row[0] or not row[1]:
                    continue
                    
                date_val = row[0]
                amount_val = row[1]
                
                if isinstance(date_val, datetime):
                    try:
                        amt = float(amount_val)
                        if amt != 0:
                            transactions.append((date_val, amt))
                    except (ValueError, TypeError):
                        pass
            
            if not transactions:
                print(f"  No transactions found in {sheet_name}")
                continue
                
            # Sort by date
            transactions.sort(key=lambda x: x[0])
            initial_date = transactions[0][0]
            
            # 1. Create Investment
            cursor.execute("""
                INSERT INTO investments (
                    institution_name, initial_investment_date, investment_type,
                    investment_name, investment_ticker, unit_currency,
                    initial_unit_price, unit_price, number_of_units_held,
                    total_dividends_received, total_tax_paid, total_fees_paid,
                    investment_fee, investment_status
                ) VALUES (
                    %s, %s, %s, %s, %s, %s,
                    %s, %s, %s, %s, %s, %s, %s, %s
                ) RETURNING id
            """, (
                inst, initial_date, inv_type, fund, ticker, curr,
                0, 0, 0, 0, 0, 0, 0, 'Active'
            ))
            inv_id = cursor.fetchone()[0]
            
            # Add to source meta based on ticker
            source = 'yfinance' if '.' in ticker or ticker == 'URTH' else 'profiledata'
            cursor.execute("""
                INSERT INTO investment_source_meta (investment_id, source, source_ticker)
                VALUES (%s, %s, %s)
            """, (inv_id, source, ticker))
            
            # 2. Add Transactions
            for date_val, amount in transactions:
                t_type = "Buy" if amount > 0 else "Sell"
                cursor.execute("""
                    INSERT INTO transactions (
                        investment_id, transaction_date, transaction_type,
                        transaction_amount, unit_price, number_of_units
                    ) VALUES (%s, %s, %s, %s, %s, %s)
                """, (
                    inv_id, date_val, t_type, abs(amount), 0, 0
                ))
                
            print(f"  ✅ Added {fund} with {len(transactions)} transactions.")
            
        conn.commit()
        print("\n🎉 Import completed successfully!")
        
    except Exception as e:
        conn.rollback()
        print(f"❌ Error during import: {e}")
    finally:
        cursor.close()
        conn.close()

if __name__ == "__main__":
    import_excel_data()
